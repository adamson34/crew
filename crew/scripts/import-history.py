#!/usr/bin/env python3
"""
Extract Excel cost build-up data into JSON for agent interpretation.

This script reads .xlsx files and outputs structured JSON preserving
all sheet names, column headers, and row data as-is. It does NOT
interpret or map columns — that is the agent's job.

Usage:
    python3 crew/scripts/import-history.py <path-to-xlsx> [--output path.json] [--sheet name]

Requires: openpyxl (pip3 install openpyxl)
"""

import argparse
import json
import sys
from datetime import datetime, date, time
from pathlib import Path


def check_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        print(
            "Error: openpyxl is required but not installed.\n"
            "Install it with: pip3 install openpyxl",
            file=sys.stderr,
        )
        return False


def serialize_value(val):
    """Convert cell values to JSON-safe types."""
    if val is None:
        return None
    if isinstance(val, (int, float, bool)):
        return val
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, time):
        return val.isoformat()
    return str(val)


def extract_sheet(ws):
    """Extract headers and rows from a worksheet."""
    rows = []
    headers = []
    header_row_idx = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        # Skip completely empty rows
        if all(cell is None for cell in row):
            continue

        # First non-empty row is the header
        if header_row_idx is None:
            headers = [str(cell).strip() if cell is not None else f"Column_{i+1}"
                       for i, cell in enumerate(row)]
            header_row_idx = row_idx
            continue

        # Data rows — trim to header length, serialize values
        row_data = [serialize_value(cell) for cell in row[:len(headers)]]

        # Skip rows that are all None after serialization
        if all(v is None for v in row_data):
            continue

        rows.append(row_data)

    return {
        "headers": headers,
        "row_count": len(rows),
        "rows": rows,
    }


def extract_workbook(xlsx_path, sheet_name=None):
    """Extract all sheets (or a specific sheet) from an Excel workbook."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    if sheet_name and sheet_name not in wb.sheetnames:
        print(
            f"Error: Sheet '{sheet_name}' not found.\n"
            f"Available sheets: {', '.join(wb.sheetnames)}",
            file=sys.stderr,
        )
        wb.close()
        sys.exit(1)

    sheets_to_read = [sheet_name] if sheet_name else wb.sheetnames
    sheets = []

    for name in sheets_to_read:
        ws = wb[name]
        sheet_data = extract_sheet(ws)
        sheet_data["name"] = name
        sheets.append(sheet_data)

    wb.close()

    return {
        "source_file": Path(xlsx_path).name,
        "extracted_at": datetime.now().isoformat(timespec="seconds"),
        "sheets": sheets,
    }


def print_summary(data):
    """Print a human-readable summary to stdout for the agent."""
    print(f"\n{'='*60}")
    print(f"  Extracted: {data['source_file']}")
    print(f"  Timestamp: {data['extracted_at']}")
    print(f"{'='*60}\n")

    for sheet in data["sheets"]:
        print(f"  Sheet: {sheet['name']}")
        print(f"  Rows:  {sheet['row_count']}")
        print(f"  Headers:")
        for i, h in enumerate(sheet["headers"], 1):
            print(f"    {i}. {h}")

        # Show first 3 rows as preview
        if sheet["rows"]:
            print(f"  Preview (first {min(3, len(sheet['rows']))} rows):")
            for row in sheet["rows"][:3]:
                preview = [str(v)[:30] if v is not None else "—" for v in row]
                print(f"    {preview}")
        print()


def main():
    parser = argparse.ArgumentParser(
        description="Extract Excel cost build-up data into JSON for agent interpretation."
    )
    parser.add_argument("xlsx", help="Path to the .xlsx file")
    parser.add_argument(
        "--output", "-o",
        help="Output JSON path (default: same directory as xlsx, with .json extension)",
    )
    parser.add_argument(
        "--sheet", "-s",
        help="Extract only this sheet (default: all sheets)",
    )

    args = parser.parse_args()

    # Validate input
    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        print(f"Error: File not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    if xlsx_path.suffix.lower() not in (".xlsx", ".xlsm"):
        print(
            f"Error: Expected .xlsx or .xlsm file, got '{xlsx_path.suffix}'.\n"
            "If you have an .xls file, re-save it as .xlsx in Excel first.",
            file=sys.stderr,
        )
        sys.exit(1)

    if not check_openpyxl():
        sys.exit(1)

    # Extract
    data = extract_workbook(str(xlsx_path), args.sheet)

    # Determine output path
    if args.output:
        output_path = Path(args.output)
    else:
        output_path = xlsx_path.with_suffix(".json")

    # Write JSON
    with open(output_path, "w") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    # Print summary for the agent
    print_summary(data)
    print(f"  JSON written to: {output_path}\n")


if __name__ == "__main__":
    main()
